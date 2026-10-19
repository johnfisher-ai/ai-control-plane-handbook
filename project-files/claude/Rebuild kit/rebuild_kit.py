#!/usr/bin/env python3
"""Rebuild kit for the downloads on The AI Control Plane site (repository johnfisher-ai/ai-control-plane-handbook).

The HTML pages are the master. These downloads are made from them and must stay in sync with them:
  deck/slides/the-ai-control-plane-slide-NN.pptx   the executive deck, one slide per file (1-58), with the page's speaker notes
  deck/img/slide-NN.png                             the slide images on the deck pages (renders of those files)
  toolkit/downloads/ai-control-plane-assessment-templates.docx   templates T1-T13 (Part 1 blank, Part 2 filled in)
  pagefind/                                         the site search index and interface (Pagefind, pinned version below)

SITE below is a copy of the repository's working tree (the folder that holds index.html, deck/, toolkit/, handbook/).

  python3 rebuild_kit.py setup SITE                 install the render fonts and Python packages, check the tools
  python3 rebuild_kit.py check SITE [--render]      report every download that is out of sync (exit code 1 if any)
  python3 rebuild_kit.py word SITE                  rebuild the Word download from the toolkit pages
  python3 rebuild_kit.py notes SITE N... | all      rewrite slides' speaker notes from the deck pages
  python3 rebuild_kit.py edit SITE N --old TEXT --new TEXT   change text on slide N (then: render N)
  python3 rebuild_kit.py title SITE N... | all      set each file's title from the slide title on its page
  python3 rebuild_kit.py render SITE N... | all     re-render slide images from the PowerPoint files
  python3 rebuild_kit.py sizes SITE                 update the file sizes shown next to the download links
  python3 rebuild_kit.py search SITE                rebuild the search index in pagefind/ (replace the whole folder in the clone)

Needs Python 3 with python-pptx, lxml, Pillow and fontTools (setup installs them), Node with the docx package (9.x),
LibreOffice and poppler (pdftoppm), and npx for the search index. Made October 15, 2026 (updated the same day: adaptation
permissions; then site search); see "claude/Rebuild kit/README.md" in the Cowork project.
"""
import argparse, codecs, collections, glob, hashlib, html, json, os, re, shutil, subprocess, sys, tempfile, zipfile
from html.parser import HTMLParser
from urllib.parse import urljoin
from xml.sax.saxutils import escape

SITE_URL = 'https://johnfisher-ai.github.io/ai-control-plane-handbook/'
TK, HB = SITE_URL + 'toolkit/', SITE_URL + 'handbook/'
LICENSE = 'https://creativecommons.org/licenses/by-nc-nd/4.0/'
DECK_URL = SITE_URL + 'deck/'
PERMISSION_URL = DECK_URL + 'notices.html#permission'
ENDNOTES_URL = HB + 'appendix-b.html'
WB = 'ai-control-plane-assessment-templates.docx'
ALL_SLIDES = list(range(1, 59))
BANNED = codecs.decode('Nalcbvag', 'rot13')  # the word the naming rule bans (kept out of this file's own text)
HL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink'
PAGEFIND = 'pagefind@1.4.0'  # the version the published index was built with; the build is deterministic for a given version


def die(msg):
    print('ERROR:', msg)
    sys.exit(2)


def pptx_path(site, n):
    return os.path.join(site, 'deck/slides', f'the-ai-control-plane-slide-{n:02d}.pptx')


# =========================================================================== a small HTML tree
VOID = {'img', 'br', 'meta', 'link', 'input', 'hr', 'source', 'wbr', 'col'}


class Node:
    def __init__(self, tag, attrs, parent):
        self.tag, self.attrs, self.parent, self.children = tag, dict(attrs), parent, []

    @property
    def cls(self):
        return (self.attrs.get('class') or '').split()

    def elems(self):
        return [c for c in self.children if isinstance(c, Node)]

    def iter(self):
        for c in self.children:
            if isinstance(c, Node):
                yield c
                yield from c.iter()

    def find(self, tag=None, cls=None):
        return next(self.findall(tag, cls), None)

    def findall(self, tag=None, cls=None):
        return (n for n in self.iter() if (tag is None or n.tag == tag) and (cls is None or cls in n.cls))

    def text(self, skip=()):
        out = []
        for c in self.children:
            if isinstance(c, str):
                out.append(c)
            elif not any(s(c) for s in skip):
                out.append('\n' if c.tag == 'br' else c.text(skip))
        return ''.join(out)


class _Tree(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = Node('#root', [], None)
        self.cur = self.root

    def handle_starttag(self, tag, attrs):
        n = Node(tag, attrs, self.cur)
        self.cur.children.append(n)
        if tag not in VOID:
            self.cur = n

    def handle_startendtag(self, tag, attrs):
        self.cur.children.append(Node(tag, attrs, self.cur))

    def handle_endtag(self, tag):
        c = self.cur
        while c is not self.root and c.tag != tag:
            c = c.parent
        if c is not self.root:
            self.cur = c.parent

    def handle_data(self, d):
        self.cur.children.append(d)


_TREES = {}


def tree(path):
    if path not in _TREES:
        p = _Tree()
        p.feed(open(path, encoding='utf-8').read())
        p.close()
        _TREES[path] = p.root
    return _TREES[path]


def norm(s):
    return re.sub(r'\s+', ' ', s or '').strip()


# =========================================================================== Word download: content from the pages
FMT = ('b', 'i', 'prompt', 'sup', 'url', 'anchor')


def fmt_key(r):
    return tuple(bool(r.get(k)) if k not in ('url', 'anchor') else r.get(k) for k in FMT)


def merge(runs):
    out = []
    for r in runs:
        if not r.get('text'):
            continue
        r = {k: v for k, v in r.items() if k == 'text' or v}
        if out and fmt_key(out[-1]) == fmt_key(r):
            out[-1]['text'] += r['text']
        else:
            out.append(dict(r))
    return out


def run_text(runs):
    return ''.join(r['text'] for r in runs)


class Ctx:
    """One page being read: its URL, the cited endnote numbers, the figure pages for cross-references."""

    def __init__(self, site, rel, cited, figs, external_only=False):
        self.site, self.rel, self.url, self.cited, self.figs = site, rel, SITE_URL + rel, cited, figs
        self.external_only = external_only


def inline_runs(node, ctx, fmt=None):
    fmt = dict(fmt or {})
    out = []
    for ch in node.children:
        if isinstance(ch, str):
            if ch:
                out.append(dict(fmt, text=ch.replace('\n', ' ')))
            continue
        tag = ch.tag
        if tag in ('strong', 'b'):
            out += inline_runs(ch, ctx, dict(fmt, b=True))
        elif tag in ('em', 'i'):
            out += inline_runs(ch, ctx, dict(fmt, i=True))
        elif tag == 'span' and 'prompt' in ch.cls:
            out += inline_runs(ch, ctx, dict(fmt, prompt=True))
        elif tag == 'sup' and 'cite' in ch.cls:
            spec = ch.text()
            for m in re.finditer(r'(\d+)(?:–(\d+))?', spec):
                a, z = int(m.group(1)), int(m.group(2) or m.group(1))
                ctx.cited.update(range(a, z + 1))
            base = {'prompt': True} if fmt.get('prompt') else {}
            for t in re.split(r'(\d+)', spec):
                if t.isdigit():
                    out.append(dict(base, text=t, sup=True, anchor=f'en_{t}'))
                elif t:
                    out.append(dict(base, text=t, sup=True))
        elif tag == 'a':
            href = ch.attrs.get('href') or ''
            url = None
            if href.startswith('http') or (href and not ctx.external_only and not href.startswith(('mailto:', 'javascript:'))):
                url = urljoin(ctx.url, href)
            out += inline_runs(ch, ctx, dict(fmt, url=url) if url else fmt)
        elif tag == 'br':
            out.append(dict(fmt, text=' '))
        elif tag == 'img':
            continue
        else:
            out += inline_runs(ch, ctx, fmt)
    return out


# cross-references the site leaves unlinked on their own page (for example "Section 7" on the Section 7 page) still link
# in the Word file, because the Word file has no pages of its own
XREF = re.compile(r'\b(Sections?|Appendix|Appendices|Figures?)\s+((?:\d+|[A-G])(?:(?:\s*(?:,|and|to|or|–)\s*)(?:\d+|[A-G]))*)\b')


def fig_pages(folder):
    m = {}
    for f in sorted(os.listdir(folder)):
        if f.endswith('.html'):
            for n in re.findall(r'<figure id="figure-(\d+)"', open(os.path.join(folder, f), encoding='utf-8').read()):
                m[int(n)] = f
    return m


def xref_target(kind, tok, hb, figs):
    gfig, hfig = figs
    if hb:
        if kind.startswith('Section') and tok.isdigit() and 1 <= int(tok) <= 13:
            return f'{HB}section-{int(tok):02d}.html'
        if kind.startswith('Append') and tok in 'ABCDE':
            return f'{HB}appendix-{tok.lower()}.html'
        if kind.startswith('Figure') and tok.isdigit() and int(tok) in hfig:
            return f'{HB}{hfig[int(tok)]}#figure-{int(tok)}'
        return None
    if kind.startswith('Section') and tok.isdigit() and 1 <= int(tok) <= 10:
        return f'{TK}section-{int(tok):02d}.html'
    if kind.startswith('Append') and tok in 'ABCDEFG':
        return f'{TK}appendix-{tok.lower()}.html'
    if kind.startswith('Figure') and tok.isdigit() and int(tok) in gfig:
        return f'{TK}{gfig[int(tok)]}#figure-{int(tok)}'
    return None


def link_xrefs(runs, figs):
    full = run_text(runs)
    spans = []
    for m in XREF.finditer(full):
        before, after = full[max(0, m.start() - 400):m.start()], full[m.end():m.end() + 30]
        sentence = re.split(r'[.!?]\s+', before)[-1]
        hb = bool(re.search(r'handbook', sentence, re.I) and not re.search(r'handbook.*\b(toolkit|this guide)\b', sentence, re.I)
                  or re.match(r'\s*(of|in) the handbook', after, re.I))
        if re.match(r'\s*of (the|this) (client )?(deck|report|template)', after, re.I) or (m.group(1).startswith('Append') and after.startswith('. ')):
            continue
        kind, off, first = m.group(1), m.start(2), True
        for t in re.finditer(r'\d+|[A-G]', m.group(2)):
            url = xref_target(kind, t.group(0), hb, figs)
            s = m.start() if first else off + t.start()
            if url:
                spans.append((s, off + t.end(), url))
            first = False
    if not spans:
        return runs
    out, pos = [], 0
    for r in runs:
        t, s0 = r['text'], pos
        pos += len(t)
        if r.get('sup') or r.get('url') or r.get('anchor') or r.get('prompt'):
            out.append(r)
            continue
        cuts = sorted({0, len(t)} | {max(0, min(len(t), x - s0)) for s, e, _ in spans for x in (s, e)})
        for a, z in zip(cuts, cuts[1:]):
            piece = dict(r, text=t[a:z])
            for s, e, url in spans:
                if s <= s0 + a and s0 + z <= e:
                    piece['url'] = url
            out.append(piece)
    return out


def block_runs(node, ctx, xrefs=True):
    runs = inline_runs(node, ctx)
    if xrefs:
        runs = link_xrefs(merge(runs), ctx.figs)
    runs = merge(runs)
    if runs:  # no leading or trailing blanks in a paragraph
        runs[0]['text'] = runs[0]['text'].lstrip()
        runs[-1]['text'] = runs[-1]['text'].rstrip()
    return merge(runs)


def replace_in_runs(runs, old, new):
    full = run_text(runs)
    i = full.find(old)
    if i < 0:
        return runs, False
    j, out, pos, done = i + len(old), [], 0, False
    for r in runs:
        t = r['text']
        s, e = pos, pos + len(t)
        pos = e
        if e <= i or s >= j:
            out.append(r)
            continue
        if i > s:
            out.append(dict(r, text=t[:i - s]))
        if not done:
            out.append(dict({k: v for k, v in r.items() if k not in ('url', 'anchor')}, text=new))
            done = True
        if j < e:
            out.append(dict(r, text=t[j - s:]))
    return merge(out), True


def plain_lines(paras):
    return [run_text(p).strip() for p in paras] or ['']


def widths(rows, header, total=9360):
    """Column widths in twentieths of a point: room for the longest word, then shares by the amount of text."""
    n = max(len(r) for r in rows)
    score, minw = [0] * n, [700] * n
    for ri, r in enumerate(rows):
        for j, c in enumerate(r):
            lines = plain_lines(c['paras'])
            word = max((len(w) for l in lines for w in re.split(r'[\s/–-]+', l)), default=0)
            per = 112 if (header and ri == 0) or c['bold'] else 100
            minw[j] = max(minw[j], word * per + 240)
            score[j] += sum(len(l) for l in lines) if any(lines) else 24  # an empty cell is room to write
    w = [(x + 12) ** 0.7 for x in score]
    fixed, out = {}, [0] * n
    for _ in range(n + 1):
        free = [j for j in range(n) if j not in fixed]
        room = total - sum(fixed.values())
        sw = sum(w[j] for j in free) or 1
        changed = False
        for j in free:
            out[j] = room * w[j] / sw
            if out[j] < minw[j]:
                fixed[j] = minw[j]
                changed = True
        if not changed:
            break
    for j, v in fixed.items():
        out[j] = v
    if sum(out) > total:
        out = [x * total / sum(out) for x in out]
    out = [int(x) for x in out]
    out[-1] += total - sum(out)
    return out


def table_block(el, ctx):
    tbl = el.find('table')
    header = tbl.find('thead') is not None
    rows = []
    for tr in tbl.findall('tr'):
        cells = []
        for c in tr.elems():
            if c.tag not in ('td', 'th'):
                continue
            ul = c.find('ul', 'cell-list')
            if ul is not None:
                paras = [block_runs(li, ctx) for li in ul.elems() if li.tag == 'li']
            else:
                r = block_runs(c, ctx)
                paras = [r] if run_text(r).strip() else []
            cells.append({'span': int(c.attrs.get('colspan') or 1), 'paras': paras,
                          'bold': any(x.get('b') for p in paras for x in p)})
        rows.append(cells)
    return {'t': 'table', 'header': header, 'widths': widths(rows, header),
            'rows': [[{'span': c['span'], 'paras': c['paras']} for c in r] for r in rows]}


def callout_block(el, ctx):
    kids, out, i = el.elems(), [], 0
    title, kind = None, 'note'
    if kids and kids[0].tag in ('h2', 'h3', 'h4', 'h5') and 'callout-sub' not in kids[0].cls:
        title = norm(kids[0].text())
        kind = 'how' if title.startswith('How to use') else 'why' if title.startswith('Why') else 'note'
        i = 1
    for k in kids[i:]:
        if k.tag in ('ul', 'ol'):
            out.append({'t': 'list', 'ordered': k.tag == 'ol', 'items': [block_runs(li, ctx) for li in k.elems() if li.tag == 'li']})
        elif k.tag in ('h3', 'h4', 'h5', 'h6'):
            out.append({'t': 'sub', 'text': norm(k.text())})
        elif k.tag == 'p':
            out.append({'t': 'p', 'runs': block_runs(k, ctx)})
        else:
            die(f'{ctx.rel}: a callout holds a <{k.tag}>, which the Word builder does not handle')
    return {'t': 'callout', 'kind': kind, 'title': title, 'children': out}


def figure_blocks(el, ctx):
    n = int(el.attrs['id'].split('-')[1])
    img = el.find('img')
    path = os.path.normpath(os.path.join(ctx.site, os.path.dirname(ctx.rel), img.attrs['src']))
    if not os.path.exists(path):
        die(f'{ctx.rel}: Figure {n} image not found: {path}')
    w, h = int(img.attrs['width']), int(img.attrs['height'])
    runs = block_runs(el.find('figcaption'), ctx)
    runs = merge([{k: v for k, v in r.items() if not (k == 'url' and v.startswith(TK) and v.endswith(f'#figure-{n}'))} for r in runs])
    return [{'t': 'figure', 'path': path, 'w': w, 'h': h, 'alt': img.attrs.get('alt', ''), 'name': f'Figure {n}'},
            {'t': 'caption', 'runs': runs}]


def content_elems(site, rel):
    """The elements of a page's main column, after the h1 and before the previous/next links."""
    main = tree(os.path.join(site, rel)).find('main')
    if main is None:
        die(f'{rel}: no <main>')
    out, started = [], False
    for el in main.elems():
        if el.tag == 'h1':
            started = True
            continue
        if not started or el.tag == 'nav' or 'dl-note' in el.cls:
            continue
        out.append(el)
    return out


def page_blocks(site, rel, part, cited, figs):
    ctx = Ctx(site, rel, cited, figs)
    out = []
    for el in content_elems(site, rel):
        if el.tag == 'h2':
            text = norm(el.text())
            m = re.match(r'(T\d+)', text)
            out.append({'t': 'h', 'level': 2, 'text': text, 'bookmark': f'{part}_{m.group(1)}' if m else None, 'pageBreak': bool(m)})
        elif el.tag == 'h3':
            out.append({'t': 'h', 'level': 3, 'text': norm(el.text())})
        elif el.tag == 'p':
            out.append({'t': 'caption' if 'caption' in el.cls else 'p', 'runs': block_runs(el, ctx)})
        elif el.tag in ('ul', 'ol'):
            out.append({'t': 'list', 'ordered': el.tag == 'ol', 'items': [block_runs(li, ctx) for li in el.elems() if li.tag == 'li']})
        elif el.tag == 'div' and 'table-wrap' in el.cls:
            out.append(table_block(el, ctx))
        elif el.tag == 'div' and 'callout-box' in el.cls:
            out.append(callout_block(el, ctx))
        elif el.tag == 'figure':
            out += figure_blocks(el, ctx)
        else:
            die(f'{rel}: <{el.tag} class="{" ".join(el.cls)}"> is not handled; extend page_blocks()')
    return out


def text_paras(site, rel, cited, replacements=()):
    """Plain paragraphs (notices, About): external links kept, cross-reference links dropped, wording adjusted."""
    ctx = Ctx(site, rel, cited, ({}, {}), external_only=True)
    out = []
    for el in content_elems(site, rel):
        if el.tag != 'p':
            die(f'{rel}: expected paragraphs only, found <{el.tag}>')
        runs = block_runs(el, ctx, xrefs=False)
        for old, new in replacements:
            if old in run_text(runs):
                runs, _ = replace_in_runs(runs, old, new)
        out.append({'t': 'p', 'runs': runs})
    joined = ' '.join(run_text(p['runs']) for p in out)
    for old, new in replacements:
        if new not in joined:
            print(f'WARNING {rel}: the page no longer says "{old[:70]}", so the Word wording "{new[:50]}" was not applied; check the notices')
    return out


def sources(site, cited):
    s = open(os.path.join(site, 'handbook/appendix-b.html'), encoding='utf-8').read()
    ent = {}
    for m in re.finditer(r'<li id="en-(\d+)" value="\d+">(.*?)</li>', s, re.S):
        body = re.sub(r'<span class="cited"[^>]*>.*?</span>', '', m.group(2), flags=re.S)
        body = re.sub(r'<sup class="cite"[^>]*>.*?</sup>', '', body, flags=re.S)
        t = _Tree()
        t.feed(body)
        t.close()
        runs = inline_runs(t.root, Ctx(site, 'handbook/appendix-b.html', set(), ({}, {}), external_only=True))
        fixed = []
        for r in runs:  # a space before each link, as the site shows them
            if r.get('url') and fixed and not fixed[-1]['text'].endswith(' '):
                fixed.append({'text': ' '})
            fixed.append(r)
        ent[int(m.group(1))] = merge(fixed)
    missing = sorted(n for n in cited if n not in ent)
    if missing:
        die(f'endnotes cited but missing from handbook/appendix-b.html: {missing}')
    return [{'n': n, 'runs': ent[n]} for n in sorted(cited)]


NOTICE_WORDING = [
    ('This guide is part of the AI Control Plane Assessment toolkit, with the client report template and the client deck.',
     'This document is part of the AI Control Plane Assessment toolkit, with the guide Assessing an AI Control Plane Project, the client report template and the client deck.'),
    ('each source is listed in Appendix D', 'each source is listed under Sources'),
    ('This guide is general information', 'This document is general information'),
]


def build_model(site):
    cited = set()
    figs = (fig_pages(os.path.join(site, 'toolkit')), fig_pages(os.path.join(site, 'handbook')))
    s7 = page_blocks(site, 'toolkit/section-07.html', 'p1', cited, figs)
    s8 = page_blocks(site, 'toolkit/section-08.html', 'p2', cited, figs)
    p3 = page_blocks(site, 'toolkit/part-3.html', 'x', cited, figs)
    intro, ok = replace_in_runs(p3[0]['runs'], 'Part 3 fills every template', 'Part 2 fills every template')
    if not ok:
        print('WARNING toolkit/part-3.html: the first paragraph no longer starts "Part 3 fills every template"; check the Part 2 introduction')
    doc = [{'t': 'h', 'level': 1, 'text': 'About this document', 'bookmark': 'about', 'pageBreak': True},
           {'t': 'p', 'runs': [{'text': 'This document holds the thirteen templates of the AI control plane assessment, T1 to T13, ready to copy and fill in. '
                                        'Part 1 has the blank templates, each with its note on how to use it. Part 2 has the same templates filled in for '
                                        'reference use case 1, engagement launch at a global audit, tax and advisory firm, each with a note on why it is filled that way.'}]},
           {'t': 'p', 'runs': [{'text': 'The method behind the templates, the workshops and the readout are in the guide '},
                               {'text': 'Assessing an AI Control Plane Project', 'i': True},
                               {'text': ' on the website: '}, {'text': TK, 'url': TK},
                               {'text': '. Section, figure and appendix numbers in the template notes refer to that guide, or to the handbook where '
                                        'the note says so, and each one links to its page.'}]},
           {'t': 'p', 'runs': [{'text': 'Conventions. ', 'b': True}, {'text': 'Text in grey italic brackets, such as '},
                               {'text': '[Name, title]', 'prompt': True},
                               {'text': ', is a prompt to replace. Add rows as needed, and keep an owner and a date against anything the client has '
                                        'confirmed. Superscript numbers refer to the sources listed at the end, numbered as in the handbook’s endnotes. '
                                        'The firm in Part 2 and all its figures are illustrative.'}]},
           {'t': 'p', 'runs': [{'text': 'License. ', 'b': True},
                               {'text': '© 2026 John Fisher, licensed under CC BY-NC-ND 4.0, with two additional permissions: you may fill in and adapt these '
                                        'templates to run assessments, and you may adapt them for other work of your own, inside your organization or for '
                                        'clients, including in paid work, if each completed report, readout or adapted version carries the credit line given '
                                        'in the notices at the end. Neither permission covers selling the templates or publishing changed versions of them.'}]},
           {'t': 'contents'},
           {'t': 'h', 'level': 1, 'text': 'Part 1. Blank templates', 'bookmark': 'part1', 'pageBreak': True},
           {'t': 'p', 'runs': [{'text': 'Each template starts with a note on how to use it: its purpose, when to fill it, who owns it, how to fill it and what '
                                        'it feeds. The client-facing report template on the website has the same structure, so the filled templates carry '
                                        'straight into the report.'}]}]
    doc += s7
    doc += [{'t': 'h', 'level': 1, 'text': 'Part 2. Templates filled in for use case 1', 'bookmark': 'part2', 'pageBreak': True},
            {'t': 'p', 'runs': intro}]
    doc += s8
    notices = text_paras(site, 'toolkit/appendix-e.html', cited, NOTICE_WORDING)
    about = text_paras(site, 'toolkit/appendix-f.html', cited)
    doc += [{'t': 'h', 'level': 1, 'text': 'Sources', 'bookmark': 'sources', 'pageBreak': True},
            {'t': 'p', 'runs': [{'text': 'The sources cited in this document, with the numbers of the shared endnote list in '},
                                {'text': 'The AI Control Plane Handbook', 'i': True}, {'text': ' (Appendix B): '},
                                {'text': ENDNOTES_URL, 'url': ENDNOTES_URL},
                                {'text': '. Vendor-authored sources are cited for product facts, not as independent validation.'}]},
            {'t': 'sources', 'items': sources(site, cited)},
            {'t': 'h', 'level': 1, 'text': 'Notices and disclosures', 'bookmark': 'notices', 'pageBreak': True}]
    doc += notices
    doc.append({'t': 'h', 'level': 1, 'text': 'About John Fisher', 'bookmark': 'author'})
    doc += about
    contents = [('about', 'About this document', 0), ('part1', 'Part 1. Blank templates', 0)]
    contents += [(b['bookmark'], b['text'], 1) for b in s7 if b['t'] == 'h' and b.get('bookmark')]
    contents += [('part2', 'Part 2. Templates filled in for use case 1', 0)]
    contents += [(b['bookmark'], b['text'], 1) for b in s8 if b['t'] == 'h' and b.get('bookmark')]
    contents += [('sources', 'Sources', 0), ('notices', 'Notices and disclosures', 0), ('author', 'About John Fisher', 0)]
    t_count = (sum(1 for c in contents if c[0].startswith('p1_')), sum(1 for c in contents if c[0].startswith('p2_')))
    if t_count != (13, 13):
        print(f'WARNING: expected T1 to T13 in Sections 7 and 8, found {t_count[0]} and {t_count[1]} template headings')
    return {'meta': {'title': 'AI Control Plane Assessment: templates T1 to T13',
                     'subject': 'Blank templates and the worked example for reference use case 1, from the toolkit Assessing an AI Control Plane Project',
                     'creator': 'John Fisher', 'site': TK},
            'contents': [{'anchor': a, 'text': t, 'level': l} for a, t, l in contents],
            'blocks': doc}


# =========================================================================== Word download: the document (docx-js)
WORD_JS = r'''
const fs = require('fs');
const {
  Document, Packer, Paragraph, TextRun, HeadingLevel, AlignmentType, Table, TableRow, TableCell, WidthType,
  ShadingType, BorderStyle, ImageRun, ExternalHyperlink, InternalHyperlink, Bookmark, LevelFormat, Header, Footer,
  PageNumber, TabStopType, VerticalAlign, TableLayoutType, HeightRule, LineRuleType,
} = require('docx');
const model = JSON.parse(fs.readFileSync(process.argv[2], 'utf8'));
const OUT = process.argv[3];
const C = { navy: '0F172A', ink: '1E293B', body: '1F2937', muted: '64748B', sky: '0369A1', line: 'CBD5E1',
  howFill: 'EEF6FC', howLine: 'B6D4EA', whyFill: 'F1F5F9', whyLine: 'CBD5E1', link: '0369A1' };
const FONT = 'Calibri';
const PAGE_W = 9360; // 6.5 in at 1 in margins on US Letter
function textRun(r, o = {}) {
  const opts = { text: r.text, font: FONT };
  if (r.b) opts.bold = true;
  if (r.i) opts.italics = true;
  if (r.prompt) { opts.italics = true; opts.color = C.muted; }
  if (r.sup) opts.superScript = true;
  if (o.size) opts.size = o.size;
  if (o.color && !r.prompt) opts.color = o.color;
  if (o.bold) opts.bold = true;
  if (r.url || r.anchor) { opts.color = C.link; if (!r.sup) opts.underline = {}; }
  return new TextRun(opts);
}
function runs(list, o = {}) {
  const out = [];
  for (const r of list || []) {
    const t = textRun(r, o);
    if (r.url) out.push(new ExternalHyperlink({ link: r.url, children: [t] }));
    else if (r.anchor) out.push(new InternalHyperlink({ anchor: r.anchor, children: [t] }));
    else out.push(t);
  }
  return out;
}
let listInstance = 0; // each list restarts its numbering
function list(b, o = {}) {
  listInstance += 1;
  return b.items.map((it) => new Paragraph({ children: runs(it, o),
    numbering: { reference: b.ordered ? 'decimal' : 'bullets', level: 0, instance: listInstance }, spacing: { after: 60 } }));
}
const thin = (color) => ({ style: BorderStyle.SINGLE, size: 4, color });
function table(b) {
  const rows = b.rows.map((r, ri) => {
    const head = b.header && ri === 0;
    const emptyRow = r.every((c) => c.paras.length === 0);
    let col = 0;
    const cells = r.map((c) => {
      const w = b.widths.slice(col, col + c.span).reduce((a, x) => a + x, 0);
      col += c.span;
      const paras = c.paras.length ? c.paras : [[]];
      return new TableCell({
        width: { size: w, type: WidthType.DXA }, columnSpan: c.span > 1 ? c.span : undefined,
        shading: head ? { type: ShadingType.CLEAR, color: 'auto', fill: C.navy } : undefined,
        margins: { top: 60, bottom: 60, left: 100, right: 100 },
        verticalAlign: head ? VerticalAlign.BOTTOM : VerticalAlign.TOP,
        children: paras.map((p, k) => new Paragraph({
          children: runs(p, head ? { size: 18, color: 'FFFFFF', bold: true } : { size: 18 }),
          spacing: { before: 0, after: k < paras.length - 1 ? 80 : 0, line: 252 } })),
      });
    });
    return new TableRow({ children: cells, tableHeader: head, cantSplit: true,
      height: emptyRow ? { value: 400, rule: HeightRule.ATLEAST } : undefined });
  });
  return [new Table({ width: { size: PAGE_W, type: WidthType.DXA }, columnWidths: b.widths, layout: TableLayoutType.FIXED,
    borders: { top: thin(C.line), bottom: thin(C.line), left: thin(C.line), right: thin(C.line), insideHorizontal: thin(C.line), insideVertical: thin(C.line) },
    rows }), new Paragraph({ children: [], spacing: { after: 120 } })];
}
function callout(b) {
  const fill = b.kind === 'how' ? C.howFill : C.whyFill;
  const line = b.kind === 'how' ? C.howLine : C.whyLine;
  const kids = [];
  if (b.title) kids.push(new Paragraph({ children: [new TextRun({ text: b.title, bold: true, color: b.kind === 'how' ? C.sky : C.ink, font: FONT, size: 22 })], spacing: { after: 100 } }));
  for (const c of b.children) {
    if (c.t === 'sub') kids.push(new Paragraph({ children: [new TextRun({ text: c.text, bold: true, font: FONT, size: 20, color: C.ink })], spacing: { before: 120, after: 60 } }));
    else if (c.t === 'p') kids.push(new Paragraph({ children: runs(c.runs, { size: 20 }), spacing: { after: 80 } }));
    else if (c.t === 'list') kids.push(...list(c, { size: 20 }));
  }
  return [new Table({ width: { size: PAGE_W, type: WidthType.DXA }, columnWidths: [PAGE_W], layout: TableLayoutType.FIXED,
    borders: { top: thin(line), bottom: thin(line), left: thin(line), right: thin(line), insideHorizontal: thin(line), insideVertical: thin(line) },
    rows: [new TableRow({ cantSplit: true, children: [new TableCell({ width: { size: PAGE_W, type: WidthType.DXA },
      shading: { type: ShadingType.CLEAR, color: 'auto', fill }, margins: { top: 140, bottom: 140, left: 200, right: 200 }, children: kids })] })],
  }), new Paragraph({ children: [], spacing: { after: 120 } })];
}
function figure(b) {
  const maxW = 624, maxH = 760; // pixels at 96 dpi: 6.5 in x 7.9 in
  let w = maxW, h = Math.round(maxW * b.h / b.w);
  if (h > maxH) { h = maxH; w = Math.round(maxH * b.w / b.h); }
  return [new Paragraph({ alignment: AlignmentType.CENTER, keepNext: true, spacing: { before: 120, after: 60, line: 240, lineRule: LineRuleType.AUTO },
    children: [new ImageRun({ type: 'png', data: fs.readFileSync(b.path), transformation: { width: w, height: h },
      altText: { title: b.name, description: b.alt, name: b.name } })] })];
}
function heading(b) {
  const level = { 1: HeadingLevel.HEADING_1, 2: HeadingLevel.HEADING_2, 3: HeadingLevel.HEADING_3 }[b.level];
  const child = new TextRun({ text: b.text, font: FONT });
  return [new Paragraph({ heading: level, pageBreakBefore: !!b.pageBreak, keepNext: true,
    children: b.bookmark ? [new Bookmark({ id: b.bookmark, children: [child] })] : [child] })];
}
const body = [];
for (const b of model.blocks) {
  if (b.t === 'h') body.push(...heading(b));
  else if (b.t === 'p') body.push(new Paragraph({ children: runs(b.runs), spacing: { after: 140 } }));
  else if (b.t === 'list') body.push(...list(b));
  else if (b.t === 'table') body.push(...table(b));
  else if (b.t === 'callout') body.push(...callout(b));
  else if (b.t === 'figure') body.push(...figure(b));
  else if (b.t === 'caption') body.push(new Paragraph({ children: runs(b.runs, { size: 18, color: C.muted }), spacing: { after: 240 } }));
  else if (b.t === 'contents') {
    body.push(new Paragraph({ heading: HeadingLevel.HEADING_2, children: [new TextRun({ text: 'Contents', font: FONT })] }));
    for (const c of model.contents) {
      body.push(new Paragraph({ indent: { left: c.level ? 360 : 0 }, spacing: { after: c.level ? 20 : 60, before: c.level ? 0 : 80 },
        children: [new InternalHyperlink({ anchor: c.anchor, children: [new TextRun({ text: c.text, font: FONT, color: C.link, bold: !c.level, size: c.level ? 20 : 22 })] })] }));
    }
  } else if (b.t === 'sources') {
    for (const s of b.items) {
      body.push(new Paragraph({ spacing: { after: 100 }, indent: { left: 480, hanging: 480 },
        children: [new Bookmark({ id: `en_${s.n}`, children: [new TextRun({ text: `${s.n}.\t`, bold: true, font: FONT, size: 19 })] }), ...runs(s.runs, { size: 19 })],
        tabStops: [{ type: TabStopType.LEFT, position: 480 }] }));
    }
  } else throw new Error('unknown block ' + b.t);
}
const cover = [
  new Paragraph({ spacing: { before: 2400, after: 160 }, children: [new TextRun({ text: 'AI CONTROL PLANE ASSESSMENT TOOLKIT', bold: true, color: C.sky, size: 20, font: FONT, characterSpacing: 40 })] }),
  new Paragraph({ spacing: { after: 240 }, children: [new TextRun({ text: 'Assessment templates T1 to T13', bold: true, color: C.navy, size: 56, font: FONT })] }),
  new Paragraph({ spacing: { after: 600 }, children: [new TextRun({ text: 'Part 1: the blank templates. Part 2: the templates filled in for reference use case 1, engagement launch at a global audit, tax and advisory firm (illustrative).', color: C.ink, size: 26, font: FONT })] }),
  new Paragraph({ spacing: { after: 80 }, children: [new TextRun({ text: 'John Fisher', bold: true, color: C.navy, size: 24, font: FONT })] }),
  new Paragraph({ spacing: { after: 1200 }, children: [new TextRun({ text: '2026', color: C.ink, size: 24, font: FONT })] }),
  new Paragraph({ spacing: { after: 80 }, children: [new TextRun({ text: 'A companion to the guide ', color: C.muted, size: 20, font: FONT }), new TextRun({ text: 'Assessing an AI Control Plane Project', italics: true, color: C.muted, size: 20, font: FONT })] }),
  new Paragraph({ children: [new ExternalHyperlink({ link: model.meta.site, children: [new TextRun({ text: model.meta.site, color: C.link, underline: {}, size: 20, font: FONT })] })] }),
  new Paragraph({ spacing: { before: 1600 }, children: [new TextRun({ text: '© 2026 John Fisher. Licensed under CC BY-NC-ND 4.0, with additional permissions to fill in and adapt the templates; see the notices at the end. John Fisher works for Salesforce; the views here are his own and drawn from public information, and this toolkit is not affiliated with or endorsed by Salesforce.', color: C.muted, size: 17, font: FONT })] }),
];
const header = new Header({ children: [new Paragraph({ children: [new TextRun({ text: 'AI Control Plane Assessment  |  Templates T1 to T13', color: C.muted, size: 16, font: FONT })] })] });
const footer = new Footer({ children: [new Paragraph({ tabStops: [{ type: TabStopType.RIGHT, position: PAGE_W }],
  children: [new TextRun({ text: 'John Fisher, 2026  ·  CC BY-NC-ND 4.0', color: C.muted, size: 16, font: FONT }),
    new TextRun({ children: ['\tPage ', PageNumber.CURRENT], color: C.muted, size: 16, font: FONT })] })] });
const doc = new Document({
  creator: model.meta.creator, title: model.meta.title, subject: model.meta.subject, lastModifiedBy: 'John Fisher',
  description: 'CC BY-NC-ND 4.0, with the AI Control Plane Assessment toolkit’s additional permissions. ' + model.meta.site,
  styles: {
    default: { document: { run: { font: FONT, size: 21, color: C.body }, paragraph: { spacing: { line: 276, lineRule: LineRuleType.AUTO } } } },
    paragraphStyles: [
      { id: 'Heading1', name: 'Heading 1', basedOn: 'Normal', next: 'Normal', quickFormat: true,
        run: { font: FONT, size: 36, bold: true, color: C.navy }, paragraph: { spacing: { before: 240, after: 200 }, outlineLevel: 0 } },
      { id: 'Heading2', name: 'Heading 2', basedOn: 'Normal', next: 'Normal', quickFormat: true,
        run: { font: FONT, size: 28, bold: true, color: C.sky }, paragraph: { spacing: { before: 240, after: 140 }, outlineLevel: 1 } },
      { id: 'Heading3', name: 'Heading 3', basedOn: 'Normal', next: 'Normal', quickFormat: true,
        run: { font: FONT, size: 23, bold: true, color: C.ink }, paragraph: { spacing: { before: 200, after: 100 }, outlineLevel: 2 } },
    ],
  },
  numbering: { config: [
    { reference: 'bullets', levels: [{ level: 0, format: LevelFormat.BULLET, text: '•', alignment: AlignmentType.LEFT, style: { paragraph: { indent: { left: 360, hanging: 240 } } } }] },
    { reference: 'decimal', levels: [{ level: 0, format: LevelFormat.DECIMAL, text: '%1.', alignment: AlignmentType.LEFT, style: { paragraph: { indent: { left: 360, hanging: 300 } } } }] },
  ] },
  features: { updateFields: false },
  sections: [{
    properties: { titlePage: true, page: { size: { width: 12240, height: 15840 }, margin: { top: 1440, right: 1440, bottom: 1440, left: 1440, header: 708, footer: 708 } } },
    headers: { default: header, first: new Header({ children: [new Paragraph({ children: [] })] }) },
    footers: { default: footer, first: new Footer({ children: [new Paragraph({ children: [] })] }) },
    children: [...cover, ...body],
  }],
});
Packer.toBuffer(doc).then((buf) => { fs.writeFileSync(OUT, buf); });
'''


def node_env():
    env = dict(os.environ)
    paths = [p for p in env.get('NODE_PATH', '').split(':') if p]
    try:
        paths.append(subprocess.run(['npm', 'root', '-g'], capture_output=True, text=True).stdout.strip())
    except FileNotFoundError:
        pass
    paths += ['/opt/npm-tools/node_modules', '/usr/local/lib/node_modules', '/usr/lib/node_modules']
    env['NODE_PATH'] = ':'.join(p for p in paths if p)
    return env


def fix_docx_ids(path):
    """docx-js writes w:id="1" on every bookmark and id 1 on every picture: give each its own."""
    tmp = path + '.tmp'
    zin = zipfile.ZipFile(path)
    with zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            d = zin.read(it.filename)
            if it.filename == 'word/document.xml':
                s = d.decode('utf-8')
                k, open_ids = [0], []

                def bm(m):
                    if m.group(1) == 'Start':
                        k[0] += 1
                        open_ids.append(k[0])
                        return m.group(0).replace(m.group(2), f'w:id="{k[0]}"')
                    return m.group(0).replace(m.group(2), f'w:id="{open_ids.pop(0)}"')
                s = re.sub(r'<w:bookmark(Start|End) [^>]*?(w:id="\d+")[^>]*/>', bm, s)
                for pat in (r'(<wp:docPr id=")\d+(")', r'(<pic:cNvPr id=")\d+(")'):
                    c = [0]

                    def num(m):
                        c[0] += 1
                        return m.group(1) + str(c[0]) + m.group(2)
                    s = re.sub(pat, num, s)
                d = s.encode('utf-8')
            zout.writestr(it, d)
    zin.close()
    shutil.move(tmp, path)


def build_docx(site, out_path):
    model = build_model(site)
    work = tempfile.mkdtemp(prefix='acp-word-')
    mp, js = os.path.join(work, 'model.json'), os.path.join(work, 'build.js')
    json.dump(model, open(mp, 'w', encoding='utf-8'), ensure_ascii=False)
    open(js, 'w', encoding='utf-8').write(WORD_JS)
    r = subprocess.run(['node', js, mp, out_path], capture_output=True, text=True, env=node_env())
    if r.returncode or not os.path.exists(out_path):
        die('node could not build the Word file (is the docx package installed? npm install -g docx@9):\n' + r.stderr[-1500:])
    fix_docx_ids(out_path)
    shutil.rmtree(work, ignore_errors=True)
    return model


def docx_parts(path):
    z = zipfile.ZipFile(path)
    doc = z.read('word/document.xml').decode('utf-8')
    paras = []
    for p in re.findall(r'<w:p[ >].*?</w:p>', doc, re.S):
        p = re.sub(r'<w:tab/>', '\t', p)
        paras.append(html.unescape(''.join(re.findall(r'<w:t(?: [^>]*)?>([^<]*)</w:t>', p))))
    rels = z.read('word/_rels/document.xml.rels').decode('utf-8')
    links = sorted(html.unescape(t) for t in re.findall(r'Target="([^"]+)"[^>]*TargetMode="External"', rels))
    media = sorted(hashlib.md5(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/'))
    bad_ids = len(set(re.findall(r'<w:bookmarkStart [^>]*w:id="(\d+)"', doc))) != len(re.findall(r'<w:bookmarkStart ', doc))
    return {'paras': paras, 'links': links, 'media': media, 'dup_bookmark_ids': bad_ids}


def compare_docx(a, b):
    """Differences between two Word files in text, links and pictures (not in dates or ids)."""
    A, B = docx_parts(a), docx_parts(b)
    out = []
    if A['paras'] != B['paras']:
        import difflib
        d = [l for l in difflib.unified_diff(A['paras'], B['paras'], 'download', 'rebuilt', n=0, lineterm='') if l[:1] in '+-' and l[:3] not in ('+++', '---')]
        out.append(f'text differs in {len(d)} paragraph lines, for example:\n    ' + '\n    '.join(x[:160] for x in d[:8]))
    if A['links'] != B['links']:
        out.append(f'links differ: only in the download {sorted(set(A["links"]) - set(B["links"]))[:5]}; only in the rebuild {sorted(set(B["links"]) - set(A["links"]))[:5]}')
    if A['media'] != B['media']:
        out.append('figures differ (an image on Section 7 or 8 changed, or a figure was added or removed)')
    return out


def banned_hits(path):
    z = zipfile.ZipFile(path)
    return [n for n in z.namelist() if n.endswith(('.xml', '.rels')) and BANNED.lower() in z.read(n).decode('utf-8', 'ignore').lower()]


def page_count(docx):
    """Pages as LibreOffice lays the file out (Word usually shows one more or one fewer)."""
    work = tempfile.mkdtemp(prefix='acp-pages-')
    try:
        soffice(['--convert-to', 'pdf', '--outdir', work, docx])
        pdf = os.path.join(work, os.path.splitext(os.path.basename(docx))[0] + '.pdf')
        out = subprocess.run(['pdfinfo', pdf], capture_output=True, text=True).stdout
        m = re.search(r'Pages:\s+(\d+)', out)
        return int(m.group(1)) if m else None
    except Exception:
        return None
    finally:
        shutil.rmtree(work, ignore_errors=True)


# =========================================================================== executive deck: pages and PowerPoint files
def deck_sections(site):
    """Slide number -> (page file, <section> node) for every slide on the deck pages."""
    out = {}
    for f in sorted(glob.glob(os.path.join(site, 'deck', '*.html'))):
        for sec in tree(f).findall('section', 'slide'):
            m = re.fullmatch(r'slide-(\d+)', sec.attrs.get('id', ''))
            if m:
                out[int(m.group(1))] = (os.path.relpath(f, site), sec)
    return out


def html_slide(sec):
    is_sup = lambda n: n.tag == 'sup'
    is_no = lambda n: 'slide-no' in n.cls
    title = norm(sec.find('h2').text(skip=(is_sup, is_no)))
    lines = []
    body = sec.find('div', 'slide-text-body')
    for el in (body.elems() if body is not None else []):
        if el.tag == 'p':
            lines += [norm(x) for x in el.text().split('\n')]
        elif el.tag == 'div' and 'table-wrap' in el.cls:
            for tr in el.findall('tr'):
                lines.append(' | '.join(norm(c.text()) for c in tr.elems() if c.tag in ('td', 'th')))
    notes = []
    nd = sec.find('div', 'notes')
    for el in (nd.elems() if nd is not None else []):
        if el.tag == 'p':
            notes.append(('p', el.text().strip()))
        elif el.tag in ('ul', 'ol'):
            notes += [('li' if el.tag == 'ul' else 'oli', li.text().strip()) for li in el.elems() if li.tag == 'li']
    cites = any('appendix-b.html#en-' in (a.attrs.get('href') or '') for a in sec.findall('a'))
    meta = sec.find('span', 'dl-meta')
    kb = re.search(r'(\d+) KB', meta.text()) if meta is not None else None
    return {'title': title, 'lines': [l for l in lines if l], 'notes': notes, 'cites': cites, 'kb': int(kb.group(1)) if kb else None}


def read_zip(path):
    z = zipfile.ZipFile(path)
    infos = z.infolist()
    data = {i.filename: z.read(i.filename) for i in infos}
    z.close()
    return infos, data


def write_zip(path, infos, data):
    tmp = path + '.tmp'
    with zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as z:
        for i in infos:
            z.writestr(i, data[i.filename])
    shutil.move(tmp, path)


def slide_parts(data):
    slide = next(n for n in data if re.fullmatch(r'ppt/slides/slide\d+\.xml', n))
    rels = data[slide.replace('slides/', 'slides/_rels/') + '.rels'].decode('utf-8')
    m = re.search(r'Target="\.\./notesSlides/(notesSlide\d+\.xml)"', rels)
    notes = f'ppt/notesSlides/{m.group(1)}' if m else None
    return slide, notes, (notes.replace('notesSlides/', 'notesSlides/_rels/') + '.rels') if notes else None


def pptx_notes(data):
    _, notes, _ = slide_parts(data)
    x = data[notes].decode('utf-8')
    m = re.search(r'<p:ph type="body"[^>]*/>.*?<p:txBody>(.*?)</p:txBody>', x, re.S)
    paras = re.findall(r'<a:p>.*?</a:p>|<a:p/>', m.group(1), re.S) if m else []
    out = [html.unescape(''.join(re.findall(r'<a:t>([^<]*)</a:t>', p))) for p in paras]
    return [p for p in out if not p.startswith(('Credit: John Fisher, The AI Control Plane: Executive Briefing',
                                                 'Bracketed numbers in these notes, and endnote numbers on the slide'))]


def pptx_title(data):
    m = re.search(r'<dc:title>(.*?)</dc:title>', data['docProps/core.xml'].decode('utf-8'), re.S)
    return html.unescape(m.group(1)) if m else ''


def pptx_text(path, n):
    """Kicker, title and body lines of the slide, read the way the deck pages' "Slide text" panels were made."""
    from pptx import Presentation
    q = '{http://schemas.openxmlformats.org/drawingml/2006/main}'

    def para(p):
        out = []
        for r in p.runs:
            t = r.text
            if not t:
                continue
            rpr = r._r.find(q + 'rPr')
            sup = rpr is not None and rpr.get('baseline') and int(rpr.get('baseline')) > 0
            out.append(('\x00' + t.strip().replace(' ', '') + '\x01') if sup and re.fullmatch(r'[\d,–\- ]+', t.strip() or 'x') else t)
        s = ''.join(out).replace('\x01\x00', ',')
        s = re.sub(r'\s+(\x00)', r'\1', s)
        return s.replace('\v', ' ').strip()

    def shapes(shs, out):
        for sh in shs:
            if sh.shape_type == 6:
                shapes(sh.shapes, out)
                continue
            if sh.has_text_frame:
                ps = [x for x in (para(p) for p in sh.text_frame.paragraphs) if x]
                if ps:
                    out.append({'top': sh.top or 0, 'left': sh.left or 0, 'type': 'p', 'paras': ps})
            if getattr(sh, 'has_table', False) and sh.has_table:
                rows = [[' '.join(x for x in (para(p) for p in c.text_frame.paragraphs) if x) for c in r.cells] for r in sh.table.rows]
                out.append({'top': sh.top or 0, 'left': sh.left or 0, 'type': 'table', 'rows': rows})

    prs = Presentation(path)
    H = prs.slide_height
    s = prs.slides[0]
    items = []
    shapes(s.shapes, items)
    items = [it for it in items if not (it['type'] == 'p' and it['paras'] == [str(n)] and it['top'] > H * 0.85)]
    srt = sorted(items, key=lambda x: (round(x['top'] / 300000), x['left']))
    pitems = [it for it in srt if it['type'] == 'p']
    isup = lambda t: re.sub(r'[^A-Za-z]', '', t).isupper()
    kit = pitems[0] if pitems and isup(pitems[0]['paras'][0]) else None
    band = round(kit['top'] / 300000) if kit else -1
    tit = next((it for it in pitems if it is not kit and round(it['top'] / 300000) > band), None) if kit else (pitems[0] if pitems else None)
    title = re.sub(r'\x00[^\x01]*\x01', '', tit['paras'][0]).strip() if tit else ''
    lines, lead = [], None
    for it in items:
        if it['type'] == 'table':
            lines += [' | '.join(c for c in row) for row in it['rows']]
            continue
        paras = it['paras'][1:] if (it is kit or it is tit) else it['paras']
        if not paras:
            continue
        if len(paras) == 1 and re.fullmatch(r'\d{1,2}', paras[0]):
            lead = paras[0]
            continue
        if lead:
            paras = [lead + '. ' + paras[0]] + paras[1:]
            lead = None
        lines += paras
    clean = lambda t: norm(t.replace('\x00', '').replace('\x01', ''))
    return {'title': norm(title), 'lines': [clean(l) for l in lines if clean(l)]}


def notes_xml(n, notes, cites, add_rel):
    def run(text, rid=None):
        link = f'<a:hlinkClick r:id="{rid}"/>' if rid else ''
        return f'<a:r><a:rPr lang="en-US" dirty="0">{link}</a:rPr><a:t>{escape(text)}</a:t></a:r>'

    def para(runs, kind='p', k=1):
        if kind == 'li':
            ppr = '<a:pPr marL="228600" indent="-228600"><a:spcAft><a:spcPts val="300"/></a:spcAft><a:buFont typeface="Arial"/><a:buChar char="•"/></a:pPr>'
        elif kind == 'oli':
            ppr = f'<a:pPr marL="228600" indent="-228600"><a:spcAft><a:spcPts val="300"/></a:spcAft><a:buFont typeface="+mj-lt"/><a:buAutoNum type="arabicPeriod" startAt="{k}"/></a:pPr>'
        else:
            ppr = '<a:pPr><a:spcAft><a:spcPts val="600"/></a:spcAft></a:pPr>'
        return f'<a:p>{ppr}{"".join(runs)}</a:p>'
    out, k = [], 0
    for kind, text in notes:
        k = k + 1 if kind == 'oli' else 0
        out.append(para([run(text)], kind, k))
    out.append(para([run(f'Credit: John Fisher, The AI Control Plane: Executive Briefing (2026), slide {n}, licensed under CC BY-NC-ND 4.0: '),
                     run(LICENSE, add_rel(LICENSE)), run('. The full deck, with its speaker notes and notices: '), run(DECK_URL, add_rel(DECK_URL)),
                     run('. You may adapt this slide, with credit, under the additional permission in the deck’s notices: '),
                     run(PERMISSION_URL, add_rel(PERMISSION_URL)), run('.')]))
    if cites:
        out.append(para([run('Bracketed numbers in these notes, and endnote numbers on the slide, refer to the endnotes in Appendix B of The AI Control Plane Handbook: '),
                         run(ENDNOTES_URL, add_rel(ENDNOTES_URL)), run('.')]))
    return ''.join(out)


def write_notes(site, n, info):
    path = pptx_path(site, n)
    infos, data = read_zip(path)
    _, notes, nrels = slide_parts(data)
    x, rels = data[notes].decode('utf-8'), data[nrels].decode('utf-8')
    existing = {html.unescape(t): i for i, t in re.findall(r'<Relationship Id="(rId\d+)" Type="' + re.escape(HL) + r'" Target="([^"]+)" TargetMode="External"/>', rels)}
    used = [int(i) for i in re.findall(r'Id="rId(\d+)"', rels)]
    new = []

    def add_rel(url):
        if url in existing:
            return existing[url]
        rid = f'rId{max(used + [0]) + 1 + len(new)}'
        existing[url] = rid
        new.append(f'<Relationship Id="{rid}" Type="{HL}" Target="{escape(url)}" TargetMode="External"/>')
        return rid
    body = notes_xml(n, info['notes'], info['cites'], add_rel)
    m = re.search(r'(<p:ph type="body"[^>]*/>.*?<p:txBody><a:bodyPr/><a:lstStyle/>)(.*?)(</p:txBody>)', x, re.S)
    if not m:
        die(f'slide {n}: the notes text box was not found')
    x = x[:m.start(2)] + body + x[m.end(2):]
    rels = rels.replace('</Relationships>', ''.join(new) + '</Relationships>')
    # drop hyperlink relationships the notes no longer use
    for rid, target in re.findall(r'<Relationship Id="(rId\d+)" Type="' + re.escape(HL) + r'" Target="([^"]+)" TargetMode="External"/>', rels):
        if f'r:id="{rid}"' not in x:
            rels = re.sub(r'<Relationship Id="' + rid + r'" Type="[^"]+" Target="[^"]+" TargetMode="External"/>', '', rels)
    changed = data[notes] != x.encode('utf-8') or data[nrels] != rels.encode('utf-8')
    data[notes], data[nrels] = x.encode('utf-8'), rels.encode('utf-8')
    if changed:
        write_zip(path, infos, data)
    return changed


def write_title(site, n, title):
    path = pptx_path(site, n)
    infos, data = read_zip(path)
    c, a = data['docProps/core.xml'].decode('utf-8'), data['docProps/app.xml'].decode('utf-8')
    c2 = re.sub(r'<dc:title>.*?</dc:title>', f'<dc:title>{escape(f"The AI Control Plane, slide {n}: {title}")}</dc:title>', c, flags=re.S)
    m = re.search(r'(<TitlesOfParts>\s*<vt:vector size="\d+" baseType="lpstr">.*)(<vt:lpstr>[^<]*</vt:lpstr>)(</vt:vector>)', a, re.S)
    a2 = a[:m.start(2)] + f'<vt:lpstr>{escape(title)}</vt:lpstr>' + a[m.end(2):] if m else a
    if (c2, a2) == (c, a):
        return False
    data['docProps/core.xml'], data['docProps/app.xml'] = c2.encode('utf-8'), a2.encode('utf-8')
    write_zip(path, infos, data)
    return True


def dup_ppr(data):
    """Paragraphs that start with two paragraph-property elements: PowerPoint asks to repair such files."""
    slide, _, _ = slide_parts(data)
    x = data[slide].decode('utf-8')
    return len(re.findall(r'<a:p>\s*(?:<a:pPr\b[^>]*/>|<a:pPr\b[^>]*>(?:(?!</a:pPr>).)*</a:pPr>)\s*<a:pPr\b', x, re.S))


def fix_dup_ppr(xml):
    pat = re.compile(r'(<a:p>\s*(?:<a:pPr\b[^>]*/>|<a:pPr\b[^>]*>(?:(?!</a:pPr>).)*</a:pPr>))\s*(?:<a:pPr\b[^>]*/>|<a:pPr\b[^>]*>(?:(?!</a:pPr>).)*</a:pPr>)', re.S)
    n = 0
    while True:
        xml, k = pat.subn(r'\1', xml)
        n += k
        if not k:
            return xml, n


def edit_slide(site, n, old, new, every=False):
    path = pptx_path(site, n)
    infos, data = read_zip(path)
    slide, _, _ = slide_parts(data)
    x = data[slide].decode('utf-8')
    run_re = re.compile(r'(<a:r>(?:(?!</a:r>).)*?<a:t>)([^<]*)(</a:t></a:r>)', re.S)
    hits = 0

    def fix_para(pm):
        nonlocal hits
        p = pm.group(0)
        runs = list(run_re.finditer(p))
        texts = [html.unescape(r.group(2)) for r in runs]
        full = ''.join(texts)
        starts = [m.start() for m in re.finditer(re.escape(old), full)]
        if not starts:
            return p
        # change only the characters that differ, so text kept from each run keeps that run's formatting
        pre = len(os.path.commonprefix([old, new]))
        suf = len(os.path.commonprefix([old[pre:][::-1], new[pre:][::-1]]))
        mid = new[pre:len(new) - suf]
        for i0 in reversed(starts):  # from the end, so earlier positions stay valid
            i, j = i0 + pre, i0 + len(old) - suf
            bounds, pos = [], 0
            for t in texts:
                bounds.append((pos, pos + len(t)))
                pos += len(t)
            if i == j:  # a pure insertion goes into the run that holds the character before it
                k = next((k for k, (s, e) in enumerate(bounds) if s < i <= e), 0)
                s = bounds[k][0]
                texts[k] = texts[k][:i - s] + mid + texts[k][i - s:]
            else:
                done = False
                for k, (s, e) in enumerate(bounds):
                    if e <= i or s >= j:
                        continue
                    t = texts[k]
                    texts[k] = t[:max(0, i - s)] + ('' if done else mid) + (t[j - s:] if j < e else '')
                    done = True
            hits += 1
        out, last = [], 0
        for r, t in zip(runs, texts):
            out.append(p[last:r.start()])
            if t or r is runs[0]:
                out.append(r.group(1) + escape(t) + r.group(3))
            last = r.end()
        out.append(p[last:])
        return ''.join(out)
    occurrences = sum(''.join(html.unescape(t) for t in re.findall(r'<a:t>([^<]*)</a:t>', pm.group(0))).count(old)
                      for pm in re.finditer(r'<a:p(?:\s[^>]*)?>.*?</a:p>', x, re.S))
    if occurrences == 0:
        die(f'slide {n}: "{old}" is not on the slide (text split across paragraphs or text boxes cannot be matched)')
    if occurrences > 1 and not every:
        die(f'slide {n}: "{old}" appears {occurrences} times; make it longer, or pass --all')
    x = re.sub(r'<a:p(?:\s[^>]*)?>.*?</a:p>', fix_para, x, flags=re.S)
    x, _ = fix_dup_ppr(x)
    data[slide] = x.encode('utf-8')
    write_zip(path, infos, data)
    return hits


# =========================================================================== rendering
def soffice(args):
    prof = tempfile.mkdtemp(prefix='acp-lo-')
    try:
        r = subprocess.run(['soffice', f'-env:UserInstallation=file://{prof}', '--headless', *args], capture_output=True, text=True, timeout=300)
        return r
    finally:
        shutil.rmtree(prof, ignore_errors=True)


def render_png(pptx, png_out):
    """Slide image as on the site: LibreOffice PDF, 120 dpi (1601 x 900), 256 colours."""
    from PIL import Image
    work = tempfile.mkdtemp(prefix='acp-render-')
    try:
        src = os.path.join(work, 's.pptx')
        shutil.copyfile(pptx, src)
        soffice(['--convert-to', 'pdf', '--outdir', work, src])
        pdf = os.path.join(work, 's.pdf')
        if not os.path.exists(pdf):
            die('LibreOffice did not convert ' + pptx)
        subprocess.run(['pdftoppm', '-png', '-r', '120', '-singlefile', pdf, os.path.join(work, 'r')], check=True)
        im = Image.open(os.path.join(work, 'r.png')).convert('RGB')
        if png_out:
            im.quantize(colors=256, method=Image.Quantize.MEDIANCUT, dither=Image.Dither.NONE).save(png_out, optimize=True)
        return im
    finally:
        shutil.rmtree(work, ignore_errors=True)


def image_differs(im, png):
    from PIL import Image, ImageChops
    cur = Image.open(png).convert('RGB')
    if cur.size != im.size:
        return f'size {cur.size} vs {im.size}'
    d = ImageChops.difference(im, cur).convert('L')
    big = sum(d.histogram()[150:])  # anti-aliasing noise stays well under 100 such pixels; one changed digit gives about 2,000
    return f'{big} pixels differ strongly' if big > 300 else None


def fonts_ok():
    out = subprocess.run(['fc-list', ':family=Plus Jakarta Sans', 'style'], capture_output=True, text=True).stdout
    return all(s in out for s in ('Regular', 'Bold', 'Italic'))


# =========================================================================== sizes shown on the pages
def update_sizes(site, pages=None):
    changed = []
    secs = deck_sections(site)
    by_page = collections.defaultdict(list)
    for n, (rel, _) in secs.items():
        by_page[rel].append(n)
    for rel, nums in by_page.items():
        p = os.path.join(site, rel)
        s = open(p, encoding='utf-8').read()
        s2 = s
        for n in nums:
            f = pptx_path(site, n)
            if not os.path.exists(f):
                continue
            kb = max(1, round(os.path.getsize(f) / 1024))
            s2 = re.sub(r'(slides/the-ai-control-plane-slide-' + f'{n:02d}' + r'\.pptx" download>[^<]*</a> <span class="dl-meta">\(\.pptx, )\d+( KB)',
                        lambda m: m.group(1) + str(kb) + m.group(2), s2)
        if s2 != s:
            open(p, 'w', encoding='utf-8').write(s2)
            changed.append(rel)
    docx = os.path.join(site, 'toolkit/downloads', WB)
    if os.path.exists(docx):
        mb = f'{os.path.getsize(docx) / 1e6:.1f} MB'
        for p in sorted(glob.glob(os.path.join(site, 'toolkit', '*.html'))):
            s = open(p, encoding='utf-8').read()
            s2 = re.sub(r'(\(\.docx, |Word document \(\.docx\), (?:US Letter, )?about \d+ pages, )\d+(?:\.\d)? MB', lambda m: m.group(1) + mb, s)
            if pages:
                s2 = re.sub(r'(Word document \(\.docx\), (?:US Letter, )?about )\d+( pages)', lambda m: m.group(1) + str(pages) + m.group(2), s2)
            if s2 != s:
                open(p, 'w', encoding='utf-8').write(s2)
                changed.append(os.path.relpath(p, site))
    return changed


# =========================================================================== commands
def parse_nums(vals):
    if vals == ['all']:
        return ALL_SLIDES
    try:
        return [int(v) for v in vals]
    except ValueError:
        die('give slide numbers (1-58) or "all"')


def cmd_setup(a):
    need = []
    for mod, pkg in (('pptx', 'python-pptx'), ('lxml', 'lxml'), ('PIL', 'Pillow'), ('fontTools', 'fonttools'), ('brotli', 'brotli')):
        try:
            __import__(mod)
        except ImportError:
            need.append(pkg)
    if need:
        print('installing', ' '.join(need))
        subprocess.run([sys.executable, '-m', 'pip', 'install', '-q', '--break-system-packages', *need], check=True)
    from fontTools.ttLib import TTFont
    dst = os.path.expanduser('~/.local/share/fonts/acp-kit')
    os.makedirs(dst, exist_ok=True)
    for w in ('400-normal', '400-italic', '700-normal'):
        src = os.path.join(a.site, 'assets/fonts', f'plus-jakarta-sans-latin-{w}.woff2')
        f = TTFont(src)
        f.flavor = None
        if 'GSUB' in f:  # no glyph substitutions (ligatures, contextual forms): the slide images were rendered without them
            del f['GSUB']
        f.save(os.path.join(dst, f'plus-jakarta-sans-{w}.ttf'))
    subprocess.run(['fc-cache', '-f'], capture_output=True)
    print('fonts:', 'Plus Jakarta Sans Regular, Bold and Italic installed' if fonts_ok() else 'NOT FOUND by fontconfig')
    for tool in ('soffice', 'pdftoppm', 'pdfinfo', 'node', 'npx'):
        print(f'{tool}:', shutil.which(tool) or 'MISSING')
    r = subprocess.run(['node', '-e', "console.log(require.resolve('docx'))"], capture_output=True, text=True, env=node_env())
    print('docx (npm):', r.stdout.strip() or 'MISSING: run npm install -g docx@9')


def cmd_check(a):
    site, problems = a.site, []
    secs = deck_sections(site)
    missing = [n for n in ALL_SLIDES if n not in secs]
    if missing:
        problems.append(f'deck pages: no section for slides {missing}')
    have_pptx = True
    try:
        import pptx  # noqa: F401
    except ImportError:
        have_pptx = False
        print('note: python-pptx is not installed, so slide text is not compared (run setup)')
    if a.render and not fonts_ok():
        die('the render fonts are not installed: run setup first')
    for n in ALL_SLIDES:
        if n not in secs:
            continue
        rel, sec = secs[n]
        f = pptx_path(site, n)
        if not os.path.exists(f):
            problems.append(f'slide {n}: {os.path.relpath(f, site)} is missing')
            continue
        info = html_slide(sec)
        _, data = read_zip(f)
        issues = []
        if [norm(t) for t in pptx_notes(data)] != [norm(t) for _, t in info['notes']]:
            issues.append('speaker notes differ from the page (run: notes %d)' % n)
        want = f'The AI Control Plane, slide {n}: {info["title"]}'
        if pptx_title(data) != want:
            issues.append(f'file title is "{pptx_title(data)}", the page says "{info["title"]}" (run: title {n})')
        if have_pptx:
            t = pptx_text(f, n)
            if t['title'] != info['title']:
                issues.append(f'slide title on the slide is "{t["title"]}", the page says "{info["title"]}" (edit the slide)')
            if t['lines'] != info['lines']:
                only_p = [l for l in t['lines'] if l not in info['lines']]
                only_h = [l for l in info['lines'] if l not in t['lines']]
                issues.append('slide text differs from the page\'s "Slide text" panel (edit the slide, then render):'
                              + ''.join(f'\n      on the slide only: {l[:150]}' for l in only_p[:4])
                              + ''.join(f'\n      on the page only:  {l[:150]}' for l in only_h[:4]))
        if dup_ppr(data):
            issues.append('paragraphs with duplicate paragraph properties: PowerPoint will ask to repair the file')
        kb = max(1, round(os.path.getsize(f) / 1024))
        if info['kb'] != kb:
            issues.append(f'the page shows {info["kb"]} KB, the file is {kb} KB (run: sizes)')
        if banned_hits(f):
            issues.append('contains the banned word: ' + ', '.join(banned_hits(f)))
        png = os.path.join(site, 'deck/img', f'slide-{n:02d}.png')
        if not os.path.exists(png):
            issues.append('slide image missing (run: render %d)' % n)
        elif a.render:
            d = image_differs(render_png(f, None), png)
            if d:
                issues.append(f'slide image does not match the PowerPoint file: {d} (run: render {n})')
        problems += [f'slide {n} ({rel}): {x}' for x in issues]
    # the Word download
    docx = os.path.join(site, 'toolkit/downloads', WB)
    if not os.path.exists(docx):
        problems.append(f'toolkit/downloads/{WB} is missing (run: word)')
    else:
        work = tempfile.mkdtemp(prefix='acp-check-')
        fresh = os.path.join(work, WB)
        build_docx(site, fresh)
        problems += [f'Word download: {d} (run: word)' for d in compare_docx(docx, fresh)]
        shutil.rmtree(work, ignore_errors=True)
        if docx_parts(docx)['dup_bookmark_ids']:
            problems.append('Word download: duplicate bookmark ids (run: word)')
        if banned_hits(docx):
            problems.append('Word download contains the banned word')
        mb = f'{os.path.getsize(docx) / 1e6:.1f} MB'
        for p in sorted(glob.glob(os.path.join(site, 'toolkit', '*.html'))):
            for shown in re.findall(r'(?:\(\.docx, |Word document \(\.docx\), (?:US Letter, )?about \d+ pages, )(\d+(?:\.\d)? MB)', open(p, encoding='utf-8').read()):
                if shown != mb:
                    problems.append(f'{os.path.relpath(p, site)} shows {shown} for the Word download, the file is {mb} (run: sizes)')
    if os.path.exists(os.path.join(site, 'search.html')):
        idx = os.path.join(site, 'pagefind')
        if not os.path.isdir(idx):
            problems.append('search index: pagefind/ is missing (run: search)')
        else:
            tmp = tempfile.mkdtemp(prefix='kit-pagefind-')
            pagefind_build(site, os.path.join(tmp, 'pagefind'))
            if folder_digest(idx) != folder_digest(os.path.join(tmp, 'pagefind')):
                problems.append('search index: pagefind/ does not match the pages (run: search)')
            shutil.rmtree(tmp, ignore_errors=True)
    print()
    if problems:
        print(f'OUT OF SYNC: {len(problems)} item(s)')
        for p in problems:
            print('  -', p)
        sys.exit(1)
    print('IN SYNC: 58 slide files, the Word download and the search index match the pages' + (' (slide images re-rendered and compared)' if a.render else ''))


def cmd_word(a):
    out = os.path.join(a.site, 'toolkit/downloads', WB)
    os.makedirs(os.path.dirname(out), exist_ok=True)
    tmp = out + '.new'
    model = build_docx(a.site, tmp)
    diffs = compare_docx(out, tmp) if os.path.exists(out) else ['new file']
    hits = banned_hits(tmp)
    if hits:
        os.remove(tmp)
        die('the rebuilt Word file contains the banned word in ' + ', '.join(hits) + '; fix the pages first')
    if not diffs:
        os.remove(tmp)
        print('Word download already matches the pages; left unchanged')
        return
    shutil.move(tmp, out)
    pages = page_count(out)
    changed = update_sizes(a.site, pages)
    n_src = next(len(b['items']) for b in model['blocks'] if b['t'] == 'sources')
    print(f'rebuilt toolkit/downloads/{WB}: {os.path.getsize(out) / 1e6:.2f} MB, {pages or "?"} pages in LibreOffice, {n_src} sources')
    for d in diffs:
        print('  changed:', d[:400])
    print('  pages updated:', ', '.join(changed) or 'none')


def cmd_notes(a):
    secs = deck_sections(a.site)
    done = [n for n in parse_nums(a.slides) if write_notes(a.site, n, html_slide(secs[n][1]))]
    print('notes rewritten:', done or 'none (already in sync)')
    if done:
        print('pages updated:', update_sizes(a.site) or 'none')


def cmd_title(a):
    secs = deck_sections(a.site)
    done = [n for n in parse_nums(a.slides) if write_title(a.site, n, html_slide(secs[n][1])['title'])]
    print('titles set:', done or 'none (already in sync)')
    if done:
        update_sizes(a.site)


def cmd_edit(a):
    k = edit_slide(a.site, a.slide, a.old, a.new, a.all)
    print(f'slide {a.slide}: {k} change(s). Now render it, look at the image, and update the page if the slide text panel or title changed.')
    update_sizes(a.site)


def cmd_render(a):
    if not fonts_ok():
        die('the render fonts are not installed: run setup first')
    for n in parse_nums(a.slides):
        png = os.path.join(a.site, 'deck/img', f'slide-{n:02d}.png')
        im = render_png(pptx_path(a.site, n), png)
        print(f'slide {n}: deck/img/slide-{n:02d}.png {im.size[0]}x{im.size[1]}')
        if im.size != (1601, 900):
            print(f'  WARNING: the pages expect 1601x900; update width and height on the <img>')


def cmd_sizes(a):
    print('pages updated:', update_sizes(a.site) or 'none (sizes already right)')


def pagefind_build(site, out):
    """Build the search index for SITE into OUT with the pinned Pagefind (npx fetches it on first use)."""
    r = subprocess.run(['npx', '-y', PAGEFIND, '--site', site, '--output-path', out], capture_output=True, text=True)
    if r.returncode != 0 or not os.path.exists(os.path.join(out, 'pagefind.js')):
        die('pagefind failed: ' + (r.stderr or r.stdout)[-800:])
    m = re.search(r'Indexed (\d+) pages', r.stdout)
    return int(m.group(1)) if m else 0


def folder_digest(folder):
    return {os.path.relpath(p, folder): hashlib.md5(open(p, 'rb').read()).hexdigest()
            for p in glob.glob(os.path.join(folder, '**', '*'), recursive=True) if os.path.isfile(p)}


def cmd_search(a):
    out = os.path.join(a.site, 'pagefind')
    tmp = tempfile.mkdtemp(prefix='kit-pagefind-')
    n = pagefind_build(a.site, os.path.join(tmp, 'pagefind'))
    new = folder_digest(os.path.join(tmp, 'pagefind'))
    if os.path.isdir(out) and folder_digest(out) == new:
        shutil.rmtree(tmp, ignore_errors=True)
        print(f'search index already matches the pages ({n} pages indexed)')
        return
    shutil.rmtree(out, ignore_errors=True)
    shutil.move(os.path.join(tmp, 'pagefind'), out)
    shutil.rmtree(tmp, ignore_errors=True)
    print(f'search index rebuilt: {n} pages, {len(new)} files. Its file names change with the content, '
          'so delete pagefind/ in the clone before unpacking the new one.')


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest='cmd', required=True)
    for name in ('setup', 'check', 'word', 'sizes', 'search'):
        p = sub.add_parser(name)
        p.add_argument('site')
        if name == 'check':
            p.add_argument('--render', action='store_true', help='also re-render every slide and compare it with its image')
    for name in ('notes', 'title', 'render'):
        p = sub.add_parser(name)
        p.add_argument('site')
        p.add_argument('slides', nargs='+')
    p = sub.add_parser('edit')
    p.add_argument('site')
    p.add_argument('slide', type=int)
    p.add_argument('--old', required=True)
    p.add_argument('--new', required=True)
    p.add_argument('--all', action='store_true', help='change every occurrence on the slide')
    a = ap.parse_args()
    a.site = os.path.abspath(a.site)
    if not os.path.exists(os.path.join(a.site, 'toolkit')) or not os.path.exists(os.path.join(a.site, 'deck')):
        die(f'{a.site} does not look like the site (no toolkit/ or deck/ folder)')
    {'setup': cmd_setup, 'check': cmd_check, 'word': cmd_word, 'notes': cmd_notes, 'title': cmd_title,
     'edit': cmd_edit, 'render': cmd_render, 'sizes': cmd_sizes, 'search': cmd_search}[a.cmd](a)


if __name__ == '__main__':
    main()
